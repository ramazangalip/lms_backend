import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.contrib.auth import get_user_model
from contents.models import Survey, SurveyQuestion, StudentSurveyResponse, PreTestResult, SurveyOption

User = get_user_model()

DEFAULT_SCALE_MAP = {
    1: "Hiçbir zaman",
    2: "Ender olarak",
    3: "Bazen",
    4: "Sıklıkla",
    5: "Her zaman"
}

def generate_survey_excel(week_number):
    # 1. Anketi ve sorularını getir
    survey = Survey.objects.filter(week_number=week_number).first()
    if not survey:
        raise ValueError(f"{week_number}. hafta için anket bulunamadı.")

    questions = list(survey.questions.all().order_by('id'))
    num_questions = len(questions)

    # 2. Ön test puanlarını, anket şık seçeneklerini ve cevaplarını getir
    pre_tests = {pt.student_id: pt.score for pt in PreTestResult.objects.all()}
    
    # Anket seçeneklerini dinamik olarak getir
    options = SurveyOption.objects.filter(question__survey=survey)
    question_options = {}
    for opt in options:
        if opt.question_id not in question_options:
            question_options[opt.question_id] = {}
        question_options[opt.question_id][opt.value] = opt.option_text

    # Cevapları hızlı arama için sözlüğe çevir (hem değeri hem de şık metnini sakla)
    responses = StudentSurveyResponse.objects.filter(question__survey=survey)
    student_responses = {}
    for resp in responses:
        if resp.student_id not in student_responses:
            student_responses[resp.student_id] = {}
        student_responses[resp.student_id][resp.question_id] = {
            'value': resp.answer_value,
            'text': resp.answer_text
        }

    # Sadece bu ankete cevap vermiş olan öğrencileri getir (cevap vermemişse Excel listesinde yer almayacak)
    answered_student_ids = list(student_responses.keys())
    students = User.objects.filter(is_student=True, id__in=answered_student_ids).order_by('department', 'first_name', 'last_name')

    # 3. Excel Çalışma Kitabını Oluştur
    wb = openpyxl.Workbook()
    # Varsayılan ilk sayfayı kaldıralım veya adını değiştirelim
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Excel stilleri
    header_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Yumuşak yeşil
    header_font = Font(name="Calibri", size=11, bold=True, color="375623")  # Koyu yeşil yazı
    data_font = Font(name="Calibri", size=11, bold=False)
    
    thin_side = Side(border_style="thin", color="D9D9D9")
    thick_bottom = Side(border_style="medium", color="375623")
    
    header_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thick_bottom)
    data_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Başlık Sütunları
    base_headers = ["Ogrenci_ID", "Ad_Soyad", "Bolum", "Sinif", "Grup", "Cinsiyet", "OnTest"]
    question_headers = [f"Soru_{i+1}" for i in range(num_questions)]
    headers = base_headers + question_headers

    def populate_sheet(ws, student_list):
        # Grid çizgilerini görünür yap
        ws.views.sheetView[0].showGridLines = True
        
        # Başlık satırını ekle ve biçimlendir
        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = header_border

        # Öğrenci satırlarını ekle
        for student in student_list:
            student_no = student.email.split('@')[0]
            
            # Sınıf hesaplama (Heuristic: 24 ile başlayanlar 2. sınıf, 25 ile başlayanlar 1. sınıf)
            sinif = ""
            if student_no.isdigit():
                if student_no.startswith("24"):
                    sinif = "2"
                elif student_no.startswith("25"):
                    sinif = "1"
                else:
                    sinif = "1"
            
            bolum = student.get_department_display() if student.department else "Belirtilmemiş"
            ontest = pre_tests.get(student.id, "")
            ad_soyad = f"{student.first_name} {student.last_name}".strip() or student_no
            
            row_data = [
                student_no,  # Ogrenci_ID
                ad_soyad,    # Ad_Soyad
                bolum,       # Bolum
                sinif,       # Sinif
                "",          # Grup (Boş bırakıyoruz)
                "",          # Cinsiyet (Boş bırakıyoruz)
                ontest       # OnTest
            ]
            
            # Soru cevaplarını ekle
            responses_dict = student_responses.get(student.id, {})
            for q in questions:
                ans_data = responses_dict.get(q.id)
                if ans_data is not None:
                    ans_val = ans_data.get('value')
                    ans_txt = ans_data.get('text')
                    
                    # 4. hafta için veri tabanındaki hatalı şık metinlerini ezmek adına her zaman varsayılan Likert eşlemesini zorunlu kılıyoruz
                    if week_number == 4:
                        final_text = DEFAULT_SCALE_MAP.get(ans_val, "")
                    else:
                        is_valid = ans_txt and str(ans_txt).strip() and str(ans_txt).lower() != 'null' and 'hata' not in str(ans_txt).lower()
                        if is_valid:
                            final_text = ans_txt
                        else:
                            final_text = question_options.get(q.id, {}).get(ans_val) or DEFAULT_SCALE_MAP.get(ans_val, "")
                        
                    if final_text:
                        # Eğer şık metni zaten sayı ile başlıyorsa (örn: "1 - Hiçbir zaman"), tekrar sayı ekleme
                        if final_text.strip() and final_text.strip()[0].isdigit():
                            cell_value = final_text
                        else:
                            cell_value = f"{ans_val} - {final_text}"
                    else:
                        cell_value = str(ans_val)
                else:
                    cell_value = ""
                row_data.append(cell_value)
            
            ws.append(row_data)
            
            # Eklenen satırı biçimlendir
            current_row = ws.max_row
            for col_num in range(1, len(row_data) + 1):
                cell = ws.cell(row=current_row, column=col_num)
                cell.font = data_font
                cell.border = data_border
                # Metinsel sütunları (Ad_Soyad ve Bolum) sola hizalayalım, diğerlerini ortalayalım
                if col_num in [2, 3]:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center

        # Sütun genişliklerini otomatik ayarla
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            # İsim sütunu (B) için daha geniş bir limit (max 45), diğer sütunlar için max 30
            if col_letter == 'B':
                ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 12)
            else:
                ws.column_dimensions[col_letter].width = max(min(max_len + 4, 30), 10)

    # A. "Tüm Öğrenciler" sayfasını oluştur
    ws_all = wb.create_sheet(title="Tüm Öğrenciler")
    populate_sheet(ws_all, students)

    # B. Bölüm bazlı sayfaları oluştur
    # Django order_by distinct hatasını önlemek için Python set kullanıyoruz
    departments = sorted(list(set(s.department for s in students if s.department)))
    for dept_code in departments:
        if not dept_code:
            continue
        # Bölümdeki öğrencileri süz
        dept_students = [s for s in students if s.department == dept_code]
        if not dept_students:
            continue
            
        # Sayfa başlığı için okunabilir ismi alalım (Excel sayfa ismi max 31 karakter olmalıdır)
        first_student = dept_students[0]
        dept_name = first_student.get_department_display() if hasattr(first_student, 'get_department_display') else dept_code
        
        # Sayfa ismini temizle ve sınırla
        safe_dept_name = str(dept_name)[:30].replace("[", "").replace("]", "").replace("*", "").replace("?", "").replace("/", "").replace("\\", "").replace(":", "")
        
        ws_dept = wb.create_sheet(title=safe_dept_name)
        populate_sheet(ws_dept, dept_students)

    # C. Soru Tanımları sayfasını oluştur (Legend)
    ws_legend = wb.create_sheet(title="Soru Açıklamaları")
    ws_legend.views.sheetView[0].showGridLines = True
    
    legend_headers = ["Soru Kodu", "Soru Metni", "Alt Boyut / Kategori"]
    ws_legend.append(legend_headers)
    
    # Legend başlık stili
    for col_num in range(1, len(legend_headers) + 1):
        cell = ws_legend.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = header_border
        
    for i, q in enumerate(questions):
        ws_legend.append([f"Soru_{i+1}", q.text, q.category or "Genel"])
        current_row = ws_legend.max_row
        for col_num in range(1, 4):
            cell = ws_legend.cell(row=current_row, column=col_num)
            cell.font = data_font
            cell.border = data_border
            if col_num == 1:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
    # Legend genişlik ayarları
    ws_legend.column_dimensions['A'].width = 15
    ws_legend.column_dimensions['B'].width = 80
    ws_legend.column_dimensions['C'].width = 30

    return wb
